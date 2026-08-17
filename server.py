import os
import json
import logging
from pathlib import Path
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse
import uvicorn
import httpx

# 環境變數
PORT = int(os.environ.get("PORT", 8080))
HOST = os.environ.get("HOST", "0.0.0.0")
WORK_DIR = Path("/app/output")
WORK_DIR.mkdir(parents=True, exist_ok=True)

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("app")

app = FastAPI(title="OfficeCLI API")


def run_cmd(cmd: str, timeout: int = 60) -> dict:
    """執行系統指令。"""
    import subprocess
    try:
        r = subprocess.run(
            cmd, shell=True, capture_output=True, text=True,
            timeout=timeout, cwd=str(WORK_DIR)
        )
        return {"ok": r.returncode == 0, "out": r.stdout.strip(), "err": r.stderr.strip()}
    except subprocess.TimeoutExpired:
        return {"ok": False, "out": "", "err": "timeout"}
    except Exception as e:
        return {"ok": False, "out": "", "err": str(e)}


@app.get("/")
def root():
    return {"status": "ok"}


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/api/create")
async def create_file(request: Request):
    """建立 Office 文件。"""
    data = await request.json()
    filename = data.get("filename", "")
    
    if not filename.endswith((".docx", ".xlsx", ".pptx")):
        return JSONResponse({"ok": False, "error": "僅支援 .docx, .xlsx, .pptx"})
    
    try:
        if filename.endswith(".docx"):
            from docx import Document
            Document().save(str(WORK_DIR / filename))
        elif filename.endswith(".xlsx"):
            from openpyxl import Workbook
            Workbook().save(str(WORK_DIR / filename))
        elif filename.endswith(".pptx"):
            from pptx import Presentation
            Presentation().save(str(WORK_DIR / filename))
        return {"ok": True, "message": f"已建立：{filename}"}
    except Exception as e:
        return JSONResponse({"ok": False, "error": str(e)})


@app.post("/api/command")
async def run_command(request: Request):
    """執行 officecli 指令。"""
    data = await request.json()
    args = data.get("args", [])
    
    if not args:
        return JSONResponse({"ok": False, "error": "請提供指令參數"})
    
    cmd = "officecli " + " ".join(args)
    result = run_cmd(cmd, timeout=120)
    return {"ok": result["ok"], "output": result["out"], "error": result["err"]}


@app.post("/api/preview")
async def preview(request: Request):
    """生成預覽圖。"""
    data = await request.json()
    filename = data.get("filename", "")
    output = data.get("output", filename.replace(".", "_") + ".png")
    
    cmd = f"officecli view {filename} screenshot -o {output}"
    result = run_cmd(cmd, timeout=120)
    return {"ok": result["ok"], "output": output, "error": result["err"]}


if __name__ == "__main__":
    logger.info("Starting server on port %d", PORT)
    uvicorn.run(app, host=HOST, port=PORT)
